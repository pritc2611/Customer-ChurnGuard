from __future__ import annotations

import io
import pandas as pd
import numpy as np
from pathlib import Path
from pydantic import BaseModel, Field, field_validator
from typing import Optional


# ─────────────────────────────────────────────────────────────────────────────
# Constants
# ─────────────────────────────────────────────────────────────────────────────
CAT_FEATURES = [
    "gender", "SeniorCitizen", "Partner", "Dependents",
    "PhoneService", "MultipleLines", "InternetService",
    "OnlineSecurity", "OnlineBackup", "DeviceProtection",
    "TechSupport", "StreamingTV", "StreamingMovies",
    "Contract", "PaperlessBilling", "PaymentMethod",
]
NUM_FEATURES = ["tenure", "MonthlyCharges", "TotalCharges"]
ALL_FEATURES = CAT_FEATURES + NUM_FEATURES

CLUSTER_LABELS = {
    0: "Loyal High-Value",
    1: "Low Engagement / Higher Risk",
}

SERVICE_COLS = [
    "OnlineSecurity", "OnlineBackup", "DeviceProtection",
    "TechSupport", "StreamingTV", "StreamingMovies",
]


# ─────────────────────────────────────────────────────────────────────────────
# HTML loader
# ─────────────────────────────────────────────────────────────────────────────
def load_html_template(filename: str) -> str:
    template_path = Path("template") / filename
    with open(template_path, "r", encoding="utf-8") as f:
        return f.read()


# ─────────────────────────────────────────────────────────────────────────────
# Pydantic input model
# ─────────────────────────────────────────────────────────────────────────────
class CustomerInput(BaseModel):
    gender: str            = Field(..., description="Male | Female")
    SeniorCitizen: str     = Field(..., description="0 | 1")
    Partner: str           = Field(..., description="Yes | No")
    Dependents: str        = Field(..., description="Yes | No")
    tenure: float          = Field(..., ge=0)
    PhoneService: str      = Field(..., description="Yes | No")
    MultipleLines: str     = Field(...)
    InternetService: str   = Field(...)
    OnlineSecurity: str    = Field(...)
    OnlineBackup: str      = Field(...)
    DeviceProtection: str  = Field(...)
    TechSupport: str       = Field(...)
    StreamingTV: str       = Field(...)
    StreamingMovies: str   = Field(...)
    Contract: str          = Field(...)
    PaperlessBilling: str  = Field(...)
    PaymentMethod: str     = Field(...)
    MonthlyCharges: float  = Field(..., ge=0)
    TotalCharges: float    = Field(..., ge=0)

    @field_validator("gender")
    def val_gender(cls, v):
        allowed = ["Male", "Female"]
        if v not in allowed:
            raise ValueError(f"gender must be one of {allowed}")
        return v

    @field_validator("SeniorCitizen")
    def val_senior(cls, v):
        if str(v) not in ["0", "1"]:
            raise ValueError("SeniorCitizen must be 0 or 1")
        return str(v)

    @field_validator("Contract")
    def val_contract(cls, v):
        allowed = ["Month-to-month", "One year", "Two year"]
        if v not in allowed:
            raise ValueError(f"Contract must be one of {allowed}")
        return v

    @field_validator("InternetService")
    def val_internet(cls, v):
        allowed = ["DSL", "Fiber optic", "No"]
        if v not in allowed:
            raise ValueError(f"InternetService must be one of {allowed}")
        return v

    def to_dataframe(self) -> pd.DataFrame:
        return pd.DataFrame([{
            "gender"           : self.gender,
            "SeniorCitizen"    : self.SeniorCitizen,
            "Partner"          : self.Partner,
            "Dependents"       : self.Dependents,
            "tenure"           : self.tenure,
            "PhoneService"     : self.PhoneService,
            "MultipleLines"    : self.MultipleLines,
            "InternetService"  : self.InternetService,
            "OnlineSecurity"   : self.OnlineSecurity,
            "OnlineBackup"     : self.OnlineBackup,
            "DeviceProtection" : self.DeviceProtection,
            "TechSupport"      : self.TechSupport,
            "StreamingTV"      : self.StreamingTV,
            "StreamingMovies"  : self.StreamingMovies,
            "Contract"         : self.Contract,
            "PaperlessBilling" : self.PaperlessBilling,
            "PaymentMethod"    : self.PaymentMethod,
            "MonthlyCharges"   : self.MonthlyCharges,
            "TotalCharges"     : self.TotalCharges,
        }])


# ─────────────────────────────────────────────────────────────────────────────
# SHAP explanation
# ─────────────────────────────────────────────────────────────────────────────
def generate_reasons(shap_val, feature_names: list[str] | None = None) -> list[str]:
    """Return human-readable SHAP factor explanations."""
    if feature_names is None:
        feature_names = ALL_FEATURES

    # shap_val.values shape: (1, n_features) for Tree or (1, n_features, 2) for PermutationExplainer
    raw = shap_val.values
    if raw.ndim == 3:
        vals = raw[0, :, 1]           # class=1 (churn)
    else:
        vals = raw[0, :]

    # Trim to actual feature count
    n = min(len(feature_names), len(vals))
    shap_df = pd.DataFrame({
        "feature"   : feature_names[:n],
        "shap_value": vals[:n],
    })
    shap_df = shap_df.reindex(
        shap_df["shap_value"].abs().sort_values(ascending=False).index
    )
    shap_df = shap_df[shap_df["shap_value"].abs() > 0.05].head(6)

    reasons = []
    for _, row in shap_df.iterrows():
        if row["shap_value"] > 0:
            reasons.append(
                f"{row['feature']} is increasing churn risk (impact: +{row['shap_value']:.3f})"
            )
        else:
            reasons.append(
                f"{row['feature']} is reducing churn risk (impact: {row['shap_value']:.3f})"
            )
    return reasons if reasons else ["No dominant factors identified."]


# ─────────────────────────────────────────────────────────────────────────────
# Segmentation
# ─────────────────────────────────────────────────────────────────────────────
def build_service_count(df: pd.DataFrame) -> pd.Series:
    return df[SERVICE_COLS].apply(lambda col: (col == "Yes").astype(int)).sum(axis=1)


def get_segment(seg_bundle: dict, df_row: pd.DataFrame) -> str:
    """Predict segment label for a single-row DataFrame."""
    scaler = seg_bundle["scaler"]
    kmeans = seg_bundle["kmeans"]

    seg_features = ["tenure", "MonthlyCharges", "TotalCharges", "ServiceCount"]
    if "ServiceCount" not in df_row.columns:
        df_row = df_row.copy()
        df_row["ServiceCount"] = build_service_count(df_row)

    scaled = scaler.transform(df_row[seg_features])
    cluster = int(kmeans.predict(scaled)[0])
    return CLUSTER_LABELS.get(cluster, "Unknown")


# ─────────────────────────────────────────────────────────────────────────────
# Risk tier
# ─────────────────────────────────────────────────────────────────────────────
def get_risk_tier(churn_probability: float) -> str:
    if churn_probability >= 0.75:
        return "Very High Risk 🚨"
    elif churn_probability >= 0.50:
        return "High Risk ⚠️"
    elif churn_probability >= 0.30:
        return "Moderate Risk ⚠️"
    else:
        return "Low Risk ✅"


def get_risk_color(churn_probability: float) -> str:
    if churn_probability >= 0.75:
        return "#dc2626"
    elif churn_probability >= 0.50:
        return "#ea580c"
    elif churn_probability >= 0.30:
        return "#d97706"
    else:
        return "#16a34a"


# ─────────────────────────────────────────────────────────────────────────────
# Business KPIs (individual)
# ─────────────────────────────────────────────────────────────────────────────
def calculate_individual_kpis(
    total_charges: float,
    monthly_charges: float,
    tenure: float,
    churn_probability: float,
    contract: str,
) -> dict:
    """Compute CLV-based business KPIs for a single customer."""
    clv = total_charges

    # CLV tier (simple thresholds instead of qcut for a single row)
    if clv < 500:
        clv_tier = "Low"
    elif clv < 2500:
        clv_tier = "Medium"
    else:
        clv_tier = "High"

    revenue_at_risk = round(clv * churn_probability, 2)

    # Expected remaining lifetime (months)
    months_remaining = (1 / churn_probability - 1) if churn_probability > 0 else 24
    expected_ltv     = round(monthly_charges * months_remaining, 2)

    # Retention cost estimate by contract
    retention_cost_map = {
        "Month-to-month": 10,
        "One year"       : 20,
        "Two year"       : 30,
    }
    retention_cost = retention_cost_map.get(contract, 10)
    net_retention_value = round(expected_ltv * 0.3 - retention_cost, 2)  # 30 % success rate

    return {
        "clv"                  : round(clv, 2),
        "clv_tier"             : clv_tier,
        "revenue_at_risk"      : revenue_at_risk,
        "expected_ltv"         : expected_ltv,
        "monthly_charges"      : round(monthly_charges, 2),
        "net_retention_value"  : net_retention_value,
    }


# ─────────────────────────────────────────────────────────────────────────────
# Batch prediction helpers
# ─────────────────────────────────────────────────────────────────────────────
def read_uploaded_file(file_bytes: bytes, filename: str) -> pd.DataFrame:
    """Read CSV or Excel bytes into a DataFrame."""
    name_lower = filename.lower()
    if name_lower.endswith(".csv"):
        return pd.read_csv(io.BytesIO(file_bytes))
    elif name_lower.endswith((".xlsx", ".xls")):
        return pd.read_excel(io.BytesIO(file_bytes))
    else:
        raise ValueError(f"Unsupported file type: {filename}. Use .csv or .xlsx")


def clean_batch_df(df: pd.DataFrame) -> pd.DataFrame:
    """Minimal cleaning for batch input."""
    df = df.copy()
    df["TotalCharges"] = pd.to_numeric(df["TotalCharges"], errors="coerce")
    df.fillna(df["MonthlyCharges"], inplace=True)
    df["SeniorCitizen"] = df["SeniorCitizen"].astype(str)
    return df


def add_engineered_features(df: pd.DataFrame) -> pd.DataFrame:
    """Add ServiceCount (needed for segmentation)."""
    df = df.copy()
    df["ServiceCount"] = build_service_count(df)
    return df


def add_kpi_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Add CLV, CLV_Tier, Revenue_At_Risk to a batch DataFrame that already has Churn_Probability."""
    df = df.copy()
    df["CLV"]             = df["TotalCharges"]
    df["Revenue_At_Risk"] = (df["CLV"] * df["Churn_Probability"]).round(2)

    # CLV tier using quantile cuts
    try:
        df["CLV_Tier"] = pd.qcut(df["CLV"], q=3, labels=["Low", "Medium", "High"])
    except Exception:
        df["CLV_Tier"] = "Medium"

    return df


def add_segment_column(df: pd.DataFrame, seg_bundle: dict) -> pd.DataFrame:
    """Vectorised segmentation for an entire DataFrame."""
    df    = df.copy()
    scaler = seg_bundle["scaler"]
    kmeans = seg_bundle["kmeans"]

    seg_features = ["tenure", "MonthlyCharges", "TotalCharges", "ServiceCount"]
    scaled  = scaler.transform(df[seg_features])
    clusters = kmeans.predict(scaled)
    df["Cluster"] = clusters
    df["Segment"] = pd.Series(clusters).map(CLUSTER_LABELS).values
    return df


def run_batch_prediction(
    df: pd.DataFrame,
    pipeline,
    seg_bundle: dict,
) -> pd.DataFrame:
    """
    Full batch inference pipeline:
      1. Clean & engineer features
      2. Predict churn probability
      3. Add segmentation
      4. Add KPIs
      5. Sort by Revenue_At_Risk descending
      6. Return enriched DataFrame
    """
    df = clean_batch_df(df)
    df = add_engineered_features(df)

    # Predict
    X              = df[ALL_FEATURES].copy()
    churn_proba    = pipeline.predict_proba(X)[:, 1]
    df["Churn_Probability"] = churn_proba.round(4)
    df["Risk_Tier"]         = [get_risk_tier(p) for p in churn_proba]

    # Segmentation
    df = add_segment_column(df, seg_bundle)

    # KPIs
    df = add_kpi_columns(df)

    # Sort by highest risk
    df = df.sort_values("Revenue_At_Risk", ascending=False).reset_index(drop=True)

    # Reorder columns for clarity
    priority_cols = [
        "Churn_Probability", "Risk_Tier", "Revenue_At_Risk",
        "CLV", "CLV_Tier", "Segment",
    ]
    id_col = ["customerID"] if "customerID" in df.columns else []
    other  = [c for c in df.columns if c not in priority_cols + id_col]
    df     = df[id_col + priority_cols + other]

    return df


def df_to_excel_bytes(df: pd.DataFrame) -> bytes:
    """Serialize DataFrame to Excel bytes with formatting."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Churn_Predictions")

        wb  = writer.book
        ws  = writer.sheets["Churn_Predictions"]

        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils import get_column_letter

        # Header style
        header_fill = PatternFill("solid", fgColor="1E3A5F")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = Alignment(horizontal="center")

        # Risk colour rows
        risk_colours = {
            "Very High Risk 🚨": "FEE2E2",
            "High Risk ⚠️"     : "FEF3C7",
            "Moderate Risk ⚠️" : "FFF7ED",
            "Low Risk ✅"      : "F0FDF4",
        }
        risk_col_idx = None
        for idx, cell in enumerate(ws[1], 1):
            if cell.value == "Risk_Tier":
                risk_col_idx = idx
                break

        if risk_col_idx:
            for row in ws.iter_rows(min_row=2):
                risk_cell = row[risk_col_idx - 1]
                colour    = risk_colours.get(str(risk_cell.value), "FFFFFF")
                fill      = PatternFill("solid", fgColor=colour)
                for cell in row:
                    cell.fill = fill

        # Auto-width columns
        for col_idx, col_cells in enumerate(ws.columns, 1):
            max_len = max((len(str(c.value)) for c in col_cells if c.value), default=10)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    return buf.getvalue()